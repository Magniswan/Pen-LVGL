from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
MD_PATH = BASE / "RK3562-GPU-NPU调研报告.md"
DOCX_PATH = BASE / "RK3562-GPU-NPU调研报告.docx"
IMAGE_PATH = BASE / "images" / "architecture.png"

BLUE = "2E74B5"
NAVY = "0B2545"
GRAY = "555555"
MUTED = "687385"
LIGHT = "F2F4F7"
CODE_BG = "F6F8FA"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_run_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name="Calibri", size=11, color=None, bold=None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def add_hyperlink(paragraph, label, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text):
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\[.*?\]\(.*?\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9, color=NAVY)
        else:
            link = re.match(r"\[(.*?)\]\((.*?)\)", token)
            if link:
                add_hyperlink(paragraph, link.group(1), link.group(2))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    set_style_font(normal, size=11, color="222222")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[style_name]
        set_style_font(style, size=size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, size=11, color="222222")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    code = doc.styles.add_style("Report Code", 1)
    set_style_font(code, name="Consolas", size=8.5, color=NAVY)
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.15)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("RK3562 GPU / RGA / NPU 调研")
    set_run_font(run, size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("LVGL Dictionary Pen PoC  ·  ")
    set_run_font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("技术调研报告")
    set_run_font(run, size=12, color=BLUE, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("RK3562 GPU / RGA / NPU 与 LVGL 加速")
    set_run_font(run, size=24, color="000000", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("面向有道 Y01 词典笔、Falcon miniapp 与全屏 LVGL 前台的可执行路线")
    set_run_font(run, size=13, color=GRAY)
    for label, value in [
        ("调研日期", "2026-08-06"),
        ("设备画像", "OVERHEAD_Y01_SKU_CHN_PRO / 固件 4.8.6 / RK3562 / AArch64"),
        ("显示基线", "LVGL 9.5.0 / DRM-KMS / 960×266 逻辑画布 / 480×960 DSI"),
        ("状态", "芯片与上游能力已确认；Y01 GPU/RGA/NPU 用户态库待真机探测"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(label + ": ")
        set_run_font(r1, size=10.5, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5)
    doc.add_paragraph()


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    widths = [int(9360 / cols)] * cols
    if cols == 4:
        widths = [1800, 2700, 2500, 2360]
    elif cols == 3:
        widths = [1800, 3600, 3960]
    elif cols == 2:
        widths = [2600, 6760]
    set_table_geometry(table, widths)
    for rindex, row in enumerate(rows):
        for cindex, value in enumerate(row):
            cell = table.cell(rindex, cindex)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, value)
            if rindex == 0:
                mark_header_row(table.rows[rindex])
                set_cell_shading(cell, LIGHT)
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body_paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    return p


def parse_markdown(doc):
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 一、"))
    lines = lines[start:]
    i = 0
    paragraph_lines = []
    def flush():
        nonlocal paragraph_lines
        if paragraph_lines:
            add_body_paragraph(doc, " ".join(x.strip() for x in paragraph_lines))
            paragraph_lines = []
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            flush()
            i += 1
            continue
        if line.startswith("!["):
            flush()
            if IMAGE_PATH.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                inline_shape = p.add_run().add_picture(str(IMAGE_PATH), width=Inches(6.35))
                inline_shape._inline.docPr.set("descr", "RK3562 GPU、RGA2、RKNPU、LVGL 与 Falcon miniapp 的分层架构图")
                inline_shape._inline.docPr.set("title", "GPU、RGA2、RKNPU 与 LVGL/Falcon 架构")
                cap = doc.add_paragraph("图 1  GPU、RGA、RKNPU 与 LVGL/Falcon 的分层架构")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    set_run_font(run, size=9, color=MUTED, italic=True)
            i += 1
            continue
        if line.startswith("## "):
            flush()
            add_body_paragraph(doc, line[3:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            flush()
            add_body_paragraph(doc, line[4:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("#### "):
            flush()
            add_body_paragraph(doc, line[5:].strip(), style="Heading 3")
            i += 1
            continue
        if line.startswith("```"):
            flush()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            p = doc.add_paragraph(style="Report Code")
            p.paragraph_format.keep_together = True
            p._p.get_or_add_pPr().append(OxmlElement("w:shd"))
            p._p.pPr[-1].set(qn("w:fill"), CODE_BG)
            run = p.add_run("\n".join(code_lines))
            set_run_font(run, name="Consolas", size=8.5, color=NAVY)
            i += 1
            continue
        if line.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                parts = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(set(x) <= set("-: ") for x in parts):
                    rows.append(parts)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("> "):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            set_run_font(run, size=10.5, color=GRAY, italic=True)
            continue
        bullet = re.match(r"^[-*] (.*)$", line)
        number = re.match(r"^\d+\. (.*)$", line)
        if bullet or number:
            flush()
            style = "List Bullet" if bullet else "List Number"
            add_body_paragraph(doc, (bullet or number).group(1), style=style)
            i += 1
            continue
        if line.strip() == "---":
            flush()
            continue
        paragraph_lines.append(line)
        i += 1
    flush()


def main():
    doc = Document()
    style_document(doc)
    add_title_block(doc)
    parse_markdown(doc)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
